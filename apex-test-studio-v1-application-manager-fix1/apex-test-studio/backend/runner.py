"""Excel-driven Timesheet UI test runner (APEX app 110543) - one browser-only video per test SAMPLE.

Prerequisites:
  pip install playwright openpyxl

Run:
  1) Start Chrome with remote debugging and log in to the app:
       & "C:\\Program Files (x86)\\Google\\Chrome\\Application\\chrome.exe" --remote-debugging-port=9222 --user-data-dir=C:\\chrome-test
  2) $env:APEX_BASE = "https://oracleapex.com/ords/r/abhi13_workspace/timesheet-app"
  3) Open the Timesheet page once in that Chrome (so the ?session=... id is visible in a tab)
  4) python apex_excel_test_runner.py

Optional environment variables:
  ALLOW_SAVE=1      allow tests that click the main Save button (TC-007, TC-008, TC-011). Default: they SKIP.
  ONLY=TC-001,TC-005  run only these test case IDs
  SLOW_MO=60        slow-motion delay in ms (default 60; 0 = fastest)
  HEADLESS=1        run without showing the browser window

Excel (timesheet_automated_test_cases.xlsx)
  Sheet Test_Cases : one row per test. TEST_DATA may hold several samples separated by ';'
                     e.g. 25;30;99  -> three samples, each with its own result row and video.
  Sheet Samples    : (optional) extra samples. TEST_CASE_ID + any column to override (TEST_DATA, DAY,
                     WEEK_OFFSET, DATE_OFFSET_DAYS, PROJECT, SAMPLE_NAME, ENABLED).
  Sheet Config     : (optional) KEY / VALUE pairs: ALLOWED_PROJECTS, DENIED_PROJECTS, SAVE_PROJECT,
                     OTHER_EMPLOYEE_ID.
"""
import os, re, csv, html, shutil, time, sys, math
from pathlib import Path
from datetime import datetime, date, timedelta
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from playwright.sync_api import sync_playwright, expect

BASE = os.environ.get("APEX_BASE", "").rstrip("/")
CDP_URL = os.environ.get("CDP_URL", "http://127.0.0.1:9222")
ALLOW_SAVE = os.environ.get("ALLOW_SAVE", "").strip().lower() in ("1", "y", "yes", "true")
HEADLESS = os.environ.get("HEADLESS", "").strip().lower() in ("1", "y", "yes", "true")
ONLY = {x.strip().upper() for x in os.environ.get("ONLY", "").split(",") if x.strip()}
EXCEL_FILE = Path(__file__).resolve().parent / "timesheet_automated_test_cases.xlsx"
OUT = Path(__file__).resolve().parent / "test_reports" / datetime.now().strftime("%Y%m%d_%H%M%S")
VIDEO_DIR = OUT / "videos"
RESULTS = []
CONFIG = {}
SESSION = ""
ON_EVENT = lambda kind, payload: None   # set by the web backend to stream progress

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
MODAL = "#ts-modal-backdrop.open"
WEEK_MIN_EDITABLE, WEEK_MAX_EDITABLE = -2, 0      # from the grid SQL: offset > 0 or < -2 is read-only


class Skip(Exception):
    """Raised by a test when it cannot run (missing precondition / feature switched off)."""


# ----------------------------------------------------------------------------- generic helpers
def settle(page, timeout=1500):
    try:
        page.wait_for_load_state("networkidle", timeout=timeout)
    except Exception:
        pass


def goto(page, alias="timesheet", extra=""):
    url = f"{BASE}/{alias}"
    sep = "?"
    if SESSION:
        url += f"?session={SESSION}"
        sep = "&"
    if extra:
        url += sep + extra
    response = page.goto(url, wait_until="load", timeout=90000)
    settle(page)
    assert response is None or response.status < 400, f"HTTP {response.status} for {alias}"
    assert "/login" not in page.url.lower(), "Redirected to login; saved browser session may have expired"
    return response


def assert_no_db_error(page):
    m = re.search(r"ORA-\d{5}[^<]{0,150}", page.content())
    assert not m, f"Database error on page: {m.group(0)}"


def page_alert_text(page):
    """Text of any visible APEX notification / error box."""
    loc = page.locator(".t-Alert--danger:visible, .t-Alert--warning:visible, #t_Alert_Notification:visible, "
                       ".apex-page-error:visible, .t-Body-alert:visible")
    out = []
    for i in range(min(loc.count(), 5)):
        try:
            out.append(loc.nth(i).inner_text().strip().replace("\n", " "))
        except Exception:
            pass
    return " | ".join(t for t in out if t)


def split_samples(value):
    if value is None:
        return [None]
    text = str(value).strip()
    if text == "":
        return [None]
    parts = [p.strip() for p in text.split(";") if p.strip() != ""]
    return parts or [None]


def to_int(value, default=0):
    try:
        if value is None or str(value).strip() == "":
            return default
        return int(float(str(value).strip()))
    except Exception:
        return default


def parse_date(text):
    text = str(text).strip()
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    if isinstance(text, str) and re.match(r"\d{4}-\d{2}-\d{2} ", text):
        return datetime.strptime(text[:10], "%Y-%m-%d").date()
    raise ValueError(f"Cannot read date '{text}' (use DD-Mon-YYYY, e.g. 02-Oct-2026)")


def monday_of(d):
    return d - timedelta(days=d.weekday())


def week_offset_for(d):
    return (monday_of(d) - monday_of(date.today())).days // 7


def fmt_hours(x):
    x = Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return format(x.normalize(), "f")


# ----------------------------------------------------------------------------- timesheet page helpers
def current_offset(page):
    try:
        return to_int(page.evaluate("() => $v('P4_WEEK_OFFSET')"), 0)
    except Exception:
        return 0


def wait_grid(page):
    page.wait_for_selector("input.timesheet-cell, .holiday-cell, .leave-cell", timeout=20000, state="attached")
    settle(page, 800)
    page.evaluate("() => { if (window.tsInitCells) tsInitCells(); }")


def open_week(page, offset=0):
    """Open the Timesheet page on the given week offset (0 = this week, -1 = last week ...)."""
    goto(page, "timesheet")
    wait_grid(page)
    cur = current_offset(page)
    guard = 0
    while cur != offset and guard < abs(offset - cur) + 4 and guard < 12:
        guard += 1
        label = "Last Week" if offset < cur else "Next Week"
        before = cur
        page.get_by_text(label, exact=True).first.click()
        page.wait_for_load_state("load")
        settle(page)
        wait_grid(page)
        cur = current_offset(page)
        if cur == before:
            break  # switcher did not move relative to the current week; use the JS fallback below
    if cur != offset:
        # Fallback: set the (hidden) item and refresh the report region, which submits P4_WEEK_OFFSET.
        page.evaluate("(o) => { apex.item('P4_WEEK_OFFSET').setValue(o);"
                      " apex.region('timesheet-log-matrix').refresh(); }", offset)
        settle(page)
        wait_grid(page)
        cur = current_offset(page)
    assert cur == offset, f"Could not navigate to week offset {offset} (page is on {cur})"
    return week_dates(page, offset)


def week_dates(page, offset=None):
    """Monday..Sunday dates of the displayed week (read from the date header row)."""
    texts = page.locator(".date-column").all_inner_texts()
    dates = []
    for t in texts:
        t = t.strip()
        if re.match(r"\d{2}-[A-Za-z]{3}-\d{4}$", t):
            dates.append(datetime.strptime(t, "%d-%b-%Y").date())
    if len(dates) >= 7:
        return dates[:7]
    base = monday_of(date.today()) + timedelta(days=7 * (offset or 0))
    return [base + timedelta(days=i) for i in range(7)]


def project_rows(page):
    """Rows that belong to a real project (have a project name)."""
    return page.locator("tr:has(.project-row-name)")


def add_first_project_if_needed(page, project=None):
    """Make sure at least one project row exists. Uses the 'Search or add a project...' box."""
    if project is None and project_rows(page).count():
        return
    names = page.locator("#project-list option").evaluate_all("els => els.map(e => e.value)")
    if project:
        match = [n for n in names if n.strip().lower() == project.strip().lower()]
        if not match and project_rows(page).filter(has_text=project).count():
            return
        assert match, f"Project '{project}' is not selectable for this employee. Selectable: {names[:10]}"
        choice = match[0]
    else:
        assert names, "No active assigned project is available for this employee"
        choice = names[0]
    box = page.locator(".project-select").first
    box.click()
    box.fill(choice)
    box.press("Enter")
    box.evaluate("el => { el.dispatchEvent(new Event('change', {bubbles:true})); el.blur(); }")
    settle(page)
    wait_grid(page)
    assert project_rows(page).count(), "Project row did not appear after choosing a project"


def cell(page, row_idx, day_idx):
    return project_rows(page).nth(row_idx).locator(f"input.timesheet-cell[name='f0{day_idx + 1}']")


def is_locked(c):
    return c.count() == 0 or c.get_attribute("data-locked") == "1" or not c.is_visible()


def cell_value(c):
    return (c.input_value() or "").strip()


def cell_is_free(page, row_idx, day_idx):
    """Editable, empty (0) and with no pending status."""
    c = cell(page, row_idx, day_idx)
    if is_locked(c):
        return False
    if cell_value(c) not in ("", "0", "0.0"):
        return False
    td = c.locator("xpath=ancestor::td[1]")
    return td.locator(".pending-status").count() == 0


def find_cell(page, free=True, days=(0, 1, 2, 3, 4, 5, 6), skip=()):
    """First (row, day) that is editable (and free). Weekdays are tried before weekend."""
    for r in range(project_rows(page).count()):
        for d in days:
            if (r, d) in skip:
                continue
            c = cell(page, r, d)
            if free:
                if cell_is_free(page, r, d):
                    return r, d
            elif not is_locked(c):
                return r, d
    return None


def open_modal(page, c):
    c.scroll_into_view_if_needed()
    c.click()
    page.wait_for_selector(MODAL, timeout=5000)


def modal_text(page):
    return page.locator("#ts-modal-error").inner_text().strip()


def close_modal(page):
    if page.locator(MODAL).count():
        try:
            page.locator("#ts-modal-backdrop .ts-btn-cancel").click(timeout=2000)
        except Exception:
            page.keyboard.press("Escape")


def apply_hours(page, c, value, desc="Automated test"):
    """Open the popup for a cell, type the hours and press Apply. Returns the popup error text ('' if accepted)."""
    open_modal(page, c)
    page.locator("#ts-modal-hours").fill(str(value))
    page.locator("#ts-modal-desc").fill(desc)
    page.locator("#ts-modal-backdrop .ts-btn-ok").click()
    page.wait_for_timeout(250)
    if page.locator(MODAL).count():
        return modal_text(page)
    return ""


def click_save(page):
    page.locator("#save-changes").click()
    page.wait_for_load_state("load")
    settle(page)
    wait_grid(page)


def require_save_allowed():
    if not ALLOW_SAVE:
        raise Skip("This test clicks the main Save button. Set $env:ALLOW_SAVE = \"1\" (test employee only) to run it.")


def cleanup_cells(page, offset, targets):
    """Reset cells to 0 and save. targets = [(row_idx, day_idx)]. Returns a short status text."""
    try:
        open_week(page, offset)
        for r, d in targets:
            c = cell(page, r, d)
            if cell_value(c) in ("", "0"):
                continue
            open_modal(page, c)
            page.locator("#ts-modal-backdrop .ts-btn-delete").click()
        click_save(page)
        left = [(r, d) for r, d in targets if cell_value(cell(page, r, d)) not in ("", "0")]
        return "cleanup OK" if not left else f"CLEANUP INCOMPLETE - remove manually: {left}"
    except Exception as exc:
        return "CLEANUP FAILED - remove test hours manually (" + str(exc).splitlines()[0][:120] + ")"


def grid_snapshot(page):
    rows = []
    for i in range(project_rows(page).count()):
        r = project_rows(page).nth(i)
        name = r.locator(".project-row-name span").first.inner_text().strip()
        vals = r.locator("input.timesheet-cell").evaluate_all("els => els.map(e => e.value)")
        rows.append((name, tuple(vals)))
    return rows


# ----------------------------------------------------------------------------- the tests
def test_hours_above_24(page, case):
    """TC-001: >24 hours is rejected in the popup (main Save is never clicked)."""
    value = case.get("TEST_DATA") or "25"
    open_week(page, to_int(case.get("WEEK_OFFSET")))
    add_first_project_if_needed(page)
    pos = find_cell(page, free=False)
    assert pos, "No editable cell found in this week"
    c = cell(page, *pos)
    before = cell_value(c)
    err = apply_hours(page, c, value, "Automated validation test: hours above 24")
    try:
        assert re.search(r"more than 24|max 24|24 hours", err, re.I), f"Expected a 24-hour validation message for '{value}', got: '{err or 'no message'}'"
        assert page.locator(MODAL).count(), "Popup closed - the invalid value was accepted"
        assert cell_value(c) == before, f"Cell value changed from {before} to {cell_value(c)}"
    finally:
        close_modal(page)
    return "PASS", f"'{value}' rejected with message: {err}. Main Save not clicked."


def test_negative_hours(page, case):
    """TC-002: negative (or otherwise non-numeric) hours are rejected."""
    value = case.get("TEST_DATA") or "-1"
    open_week(page, to_int(case.get("WEEK_OFFSET")))
    add_first_project_if_needed(page)
    pos = find_cell(page, free=False)
    assert pos, "No editable cell found in this week"
    c = cell(page, *pos)
    before = cell_value(c)
    err = apply_hours(page, c, value, "Automated validation test: negative hours")
    try:
        assert err, f"No validation message appeared for '{value}'"
        assert page.locator(MODAL).count(), "Popup closed - the invalid value was accepted"
        assert cell_value(c) == before, f"Cell value changed from {before} to {cell_value(c)}"
    finally:
        close_modal(page)
    return "PASS", f"'{value}' rejected with message: {err}"


def _assert_locked_and_not_editable(page, c, label):
    assert c.count() > 0, f"{label}: no input cell exists (treated as read-only)"
    assert c.get_attribute("data-locked") == "1", f"{label} is EDITABLE (cell is not locked)"
    c.scroll_into_view_if_needed()
    c.click()
    page.wait_for_selector(MODAL, timeout=5000)
    assert page.locator("#ts-modal-hours").get_attribute("readonly") is not None, f"{label}: popup hours box is editable"
    assert not page.locator("#ts-modal-backdrop .ts-btn-ok").is_visible(), f"{label}: Apply button is visible on a locked cell"
    close_modal(page)


def test_future_date_readonly(page, case):
    """TC-003: a date after today is read-only. DATE_OFFSET_DAYS = 1 means tomorrow."""
    target = date.today() + timedelta(days=to_int(case.get("DATE_OFFSET_DAYS"), 1))
    if target <= date.today():
        raise Skip("DATE_OFFSET_DAYS must be a positive number for a future-date test")
    dates = open_week(page, week_offset_for(target))
    add_first_project_if_needed(page)
    day = dates.index(target)
    c = cell(page, 0, day)
    _assert_locked_and_not_editable(page, c, f"{DAYS[day]} {target:%d-%b-%Y} (future)")
    return "PASS", f"{DAYS[day]} {target:%d-%b-%Y} is read-only; popup opens view-only (no Apply button)"


def test_old_week_readonly(page, case):
    """TC-004: weeks older than 2 weeks back (offset < -2) are view-only. Editable window = offsets 0, -1, -2."""
    off = to_int(case.get("WEEK_OFFSET"), -3)
    if off >= WEEK_MIN_EDITABLE:
        raise Skip(f"WEEK_OFFSET {off} is inside the editable window ({WEEK_MIN_EDITABLE}..0). Use -3 or older.")
    open_week(page, off)
    add_first_project_if_needed(page)
    n = project_rows(page).count()
    unlocked = []
    for r in range(n):
        for d in range(7):
            c = cell(page, r, d)
            if c.count() and c.get_attribute("data-locked") != "1":
                unlocked.append((r, DAYS[d]))
    assert not unlocked, f"Week {off} has editable cells: {unlocked[:6]}"
    c = cell(page, 0, 0)
    _assert_locked_and_not_editable(page, c, f"Week {off} Monday")
    return "PASS", f"All cells in week offset {off} are read-only ({n} project row(s) checked)"


def test_old_week_boundary_editable(page, case):
    """Boundary control: the oldest allowed week (offset -2) must still be editable."""
    off = to_int(case.get("WEEK_OFFSET"), WEEK_MIN_EDITABLE)
    open_week(page, off)
    add_first_project_if_needed(page)
    pos = find_cell(page, free=False)
    assert pos, f"Week {off} should be editable but every cell is locked (holiday/leave could also cause this)"
    return "PASS", f"Week offset {off} has an editable cell ({DAYS[pos[1]]})"


def test_weekend_readonly(page, case):
    """TC-005 / TC-006: Saturday / Sunday cell is read-only."""
    day_name = str(case.get("DAY") or "").strip().title()[:3]
    if day_name not in ("Sat", "Sun"):
        m = re.search(r"saturday|sunday", str(case.get("TEST_NAME", "")), re.I)
        day_name = m.group(0).title()[:3] if m else ""
    if day_name not in ("Sat", "Sun"):
        raise Skip("Add DAY = Sat or Sun in the Excel row")
    off = to_int(case.get("WEEK_OFFSET"), 0)
    note = ""
    d_idx = DAYS.index(day_name)
    if off == 0 and monday_of(date.today()) + timedelta(days=d_idx) > date.today():
        off = -1   # this week's weekend is in the future (locked for that reason) - use last week so the test is meaningful
        note = " (used last week because this weekend is still in the future)"
    dates = open_week(page, off)
    add_first_project_if_needed(page)
    if page.locator(".holiday-cell, .leave-cell").count() and not cell(page, 0, d_idx).count():
        raise Skip(f"{day_name} {dates[d_idx]:%d-%b-%Y} is a holiday/leave day - choose another WEEK_OFFSET")
    c = cell(page, 0, d_idx)
    _assert_locked_and_not_editable(page, c, f"{day_name} {dates[d_idx]:%d-%b-%Y}")
    return "PASS", f"{day_name} {dates[d_idx]:%d-%b-%Y} is read-only{note}"


def _save_one_entry(page, case, hours, label):
    require_save_allowed()
    off = to_int(case.get("WEEK_OFFSET"), 0)
    open_week(page, off)
    add_first_project_if_needed(page, CONFIG.get("SAVE_PROJECT") or None)
    pos = find_cell(page, free=True, days=(0, 1, 2, 3, 4))
    if not pos:
        # this week may have no free working day yet (e.g. Monday) - try last week
        off = off - 1 if off - 1 >= WEEK_MIN_EDITABLE else off
        open_week(page, off)
        add_first_project_if_needed(page, CONFIG.get("SAVE_PROJECT") or None)
        pos = find_cell(page, free=True, days=(0, 1, 2, 3, 4))
    assert pos, "No free editable weekday cell available for a test entry"
    c = cell(page, *pos)
    err = apply_hours(page, c, hours, f"Automated test - {label}")
    assert not err, f"Popup rejected {hours} hours: {err}"
    click_save(page)
    assert_no_db_error(page)
    alert = page_alert_text(page)
    c2 = cell(page, *pos)
    saved = cell_value(c2)
    pending = c2.locator("xpath=ancestor::td[1]").locator(".pending-status").count() > 0
    return off, pos, saved, pending, alert


def test_overtime_requires_approval(page, case):
    """TC-007: more than 8 hours -> saved entry shows PENDING."""
    hours = case.get("TEST_DATA") or "9"
    off = pos = None
    try:
        off, pos, saved, pending, alert = _save_one_entry(page, case, hours, "overtime")
        assert float(saved or 0) == float(hours), f"Saved value is '{saved}', expected {hours}. {alert}"
        assert pending, "No 'Pending' status shown next to the saved overtime hours"
        result = ("PASS", f"{hours}h saved and shown as Pending. (APPROVAL_REQUIRED=Y is not visible in the UI - confirm in the table if needed.)")
    finally:
        if pos is not None:
            msg = cleanup_cells(page, off, [pos])
    return result[0], result[1] + " " + msg


def test_normal_hours_approved(page, case):
    """TC-008: 8 hours or less is saved and is NOT pending (auto-approved)."""
    hours = case.get("TEST_DATA") or "8"
    off = pos = None
    try:
        off, pos, saved, pending, alert = _save_one_entry(page, case, hours, "normal hours")
        assert float(saved or 0) == float(hours), f"Saved value is '{saved}', expected {hours}. {alert}"
        assert not pending, "Normal hours are showing as Pending"
        result = ("PASS", f"{hours}h saved with no Pending flag (treated as approved).")
    finally:
        if pos is not None:
            msg = cleanup_cells(page, off, [pos])
    return result[0], result[1] + " " + msg


def test_assigned_active_projects_only(page, case):
    """TC-009: project drop-down lists only active projects assigned to the employee (needs Config values)."""
    open_week(page, 0)
    options = page.locator("#project-list option").evaluate_all("els => els.map(e => e.value.trim())")
    shown = [project_rows(page).nth(i).locator(".project-row-name span").first.inner_text().strip()
             for i in range(project_rows(page).count())]
    available = {x.lower() for x in options + shown}
    allowed = [x.strip() for x in re.split(r"[,|]", CONFIG.get("ALLOWED_PROJECTS", "")) if x.strip()]
    denied = [x.strip() for x in re.split(r"[,|]", CONFIG.get("DENIED_PROJECTS", "")) if x.strip()]
    if case.get("TEST_DATA"):
        denied.append(str(case["TEST_DATA"]).strip())
    if not allowed and not denied:
        raise Skip("Fill ALLOWED_PROJECTS / DENIED_PROJECTS in the Config sheet. Currently selectable: "
                   + ", ".join(sorted(set(options + shown))) or "(none)")
    missing = [p for p in allowed if p.lower() not in available]
    leaked = [p for p in denied if p.lower() in available]
    assert not missing, f"Assigned active project(s) missing from the list: {missing}"
    assert not leaked, f"Unassigned/inactive project(s) are selectable: {leaked}"
    return "PASS", f"{len(allowed)} allowed project(s) present, {len(denied)} blocked project(s) absent"


def test_cross_employee_access_denied(page, case):
    """TC-010: tampering with employee-related URL parameters never shows/changes another employee's data."""
    other = str(case.get("TEST_DATA") or CONFIG.get("OTHER_EMPLOYEE_ID") or "1").strip()
    open_week(page, 0)
    baseline = grid_snapshot(page)
    tried = []
    for param in ("P4_USER_ID", "P4_EMPLOYEE_ID", "P4_EMP_ID", "P4_EMAIL"):
        goto(page, "timesheet", f"{param}={other}&P4_WEEK_OFFSET=0")
        assert_no_db_error(page)
        text = page.content().lower()
        blocked = ("checksum" in text or "session state protection" in text or "not authorized" in text
                   or "access denied" in text or "unauthorized" in text)
        if not blocked:
            wait_grid(page)
            assert grid_snapshot(page) == baseline, f"Grid data changed when {param}={other} was supplied!"
        tried.append(param + (" (blocked)" if blocked else " (ignored)"))
    return "PASS", ("Own data unchanged after tampering: " + ", ".join(tried) +
                    ". A full test needs a 2nd login for another employee; this checks URL tampering only.")


def test_weekly_limit(page, case):
    """TC-011: a weekly total above the limit is rejected on Save."""
    require_save_allowed()
    total = Decimal(str(case.get("TEST_DATA") or "49"))
    per_day_max = Decimal("12")
    needed = int(math.ceil(total / per_day_max))
    off = to_int(case.get("WEEK_OFFSET"), 0)
    slots = []
    for candidate in dict.fromkeys([off, -1, -2]):
        if candidate < WEEK_MIN_EDITABLE:
            continue
        open_week(page, candidate)
        add_first_project_if_needed(page, CONFIG.get("SAVE_PROJECT") or None)
        slots = []
        for d in (0, 1, 2, 3, 4, 5, 6):
            if cell_is_free(page, 0, d):
                slots.append((0, d))
        if len(slots) >= needed:
            off = candidate
            break
    if len(slots) < needed:
        raise Skip(f"Need {needed} free editable days in one week to enter {total}h; found {len(slots)}")
    slots = slots[:needed]
    amounts, remaining = [], total
    for i in range(needed):
        left = needed - i
        a = (remaining / left).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        a = remaining if left == 1 else a
        amounts.append(a)
        remaining -= a
    result = None
    try:
        for (r, d), a in zip(slots, amounts):
            err = apply_hours(page, cell(page, r, d), fmt_hours(a), "Automated weekly limit test")
            assert not err, f"Popup rejected {a}h on {DAYS[d]} before Save: {err}"
        click_save(page)
        alert = page_alert_text(page)
        saved_total = Decimal(str(page.locator(".weekly-total").first.inner_text().strip() or "0"))
        saved_any = saved_total >= total
        assert not saved_any, f"Weekly total {saved_total} (>= {total}) was SAVED - no weekly limit enforced"
        result = ("PASS", f"{total}h over {needed} day(s) was not saved (week total {saved_total}). App message: {alert or 'none shown'}")
    finally:
        msg = cleanup_cells(page, off, slots)
    return result[0], result[1] + " " + msg


def test_holiday_readonly(page, case):
    """TC-012: a holiday shows HOLIDAY and cannot be edited. TEST_DATA = holiday date, or blank to auto-detect."""
    wanted = case.get("TEST_DATA")
    if wanted:
        d = parse_date(wanted)
        off = week_offset_for(d)
        if off < WEEK_MIN_EDITABLE or off > WEEK_MAX_EDITABLE or d > date.today():
            raise Skip(f"{d:%d-%b-%Y} is outside the editable window (past 3 weeks, not future); the holiday rule cannot be isolated there")
        dates = open_week(page, off)
        add_first_project_if_needed(page)
        idx = dates.index(d)
        found = [(off, idx)]
    else:
        found = []
        for off in (0, -1, -2):
            dates = open_week(page, off)
            add_first_project_if_needed(page)
            row = project_rows(page).first
            classes = row.locator(".ts-day-column").evaluate_all("els => els.map(e => e.className)")
            for j, cls in enumerate(classes[:7]):
                if "holiday-cell" in cls:
                    found.append((off, j))
            if found:
                break
        if not found:
            raise Skip("No holiday found in the last 3 weeks. Put a date from XXEA_HOLIDAYS in TEST_DATA (DD-Mon-YYYY, inside the past 3 weeks).")
        off = found[0][0]
        dates = week_dates(page, off)
    for off, idx in found[:1]:
        row = project_rows(page).first
        holiday_cells = row.locator(".holiday-cell")
        assert holiday_cells.count() >= 1, f"{dates[idx]:%d-%b-%Y} is not marked HOLIDAY on the grid"
        assert "HOLIDAY" in holiday_cells.first.inner_text().upper(), "Holiday cell has no HOLIDAY label"
        visible_inputs = holiday_cells.first.locator("input.timesheet-cell:visible")
        assert visible_inputs.count() == 0, "An editable input is visible inside the holiday cell"
    return "PASS", f"{dates[found[0][1]]:%d-%b-%Y} shows HOLIDAY and has no editable input"



# ----------------------------------------------------------------------------- custom (no-code) tests built in the web UI
STEP_LABELS = {
    "open": "Open page", "open_week": "Go to week", "click": "Click", "type": "Type into field",
    "choose": "Choose option", "press": "Press key", "wait": "Wait", "see": "Expect to see text",
    "not_see": "Expect NOT to see text", "url": "Expect URL to contain",
}


def _find(page, candidates, what, timeout=8000):
    """Return the first candidate locator that exists on the page, polling until timeout."""
    waited = 0
    while True:
        for loc in candidates:
            try:
                if loc.count() > 0:
                    return loc.first
            except Exception:
                pass
        if waited >= timeout:
            raise AssertionError(f"could not find {what} on the page")
        page.wait_for_timeout(250)
        waited += 250


def _is_css(t):
    return t.lower().startswith("css=") or bool(re.match(r"^[#.\[]", t)) or ">" in t


def _click_target(page, t):
    if _is_css(t):
        return _find(page, [page.locator(t[4:] if t.lower().startswith("css=") else t)], f"'{t}'")
    return _find(page, [page.get_by_role("button", name=t), page.get_by_role("link", name=t),
                        page.get_by_text(t, exact=True), page.get_by_text(t)], f"a button, link or text '{t}'")


def _field_target(page, t):
    if _is_css(t):
        return _find(page, [page.locator(t[4:] if t.lower().startswith("css=") else t)], f"field '{t}'")
    cands = [page.get_by_label(t), page.get_by_placeholder(t), page.locator(f"[name='{t}']")]
    if re.fullmatch(r"[A-Za-z][\w-]*", t):
        cands.append(page.locator(f"#{t}"))
    return _find(page, cands, f"field '{t}'")


def run_step(page, step):
    act = str(step.get("action") or "").strip()
    target = str(step.get("target") or "").strip()
    value = str(step.get("value") or "").strip()
    if act == "open":
        if not target:
            raise AssertionError("enter a page alias (e.g. timesheet) or a full URL")
        if target.lower().startswith("http"):
            page.goto(target, wait_until="load", timeout=90000)
            settle(page)
        else:
            goto(page, target)
    elif act == "open_week":
        open_week(page, to_int(value or target, 0))
    elif act == "click":
        _click_target(page, target).click()
        settle(page, 4000)
    elif act == "type":
        _field_target(page, target).fill(value)
    elif act == "choose":
        _field_target(page, target).select_option(label=value)
    elif act == "press":
        page.keyboard.press(value or target or "Enter")
        settle(page, 4000)
    elif act == "wait":
        page.wait_for_timeout(int(float(value or target or "1") * 1000))
    elif act == "see":
        try:
            expect(page.get_by_text(target).first).to_be_visible(timeout=8000)
        except AssertionError:
            raise AssertionError(f"the text '{target}' was not visible on the page")
    elif act == "not_see":
        try:
            expect(page.get_by_text(target)).to_have_count(0, timeout=6000)
        except AssertionError:
            raise AssertionError(f"the text '{target}' is still visible on the page")
    elif act == "url":
        needle = value or target
        assert needle in page.url, f"URL is {page.url}"
    else:
        raise AssertionError(f"unknown action '{act}'")


def test_custom_steps(page, case):
    """A test built in the web UI: runs the listed steps in order, stops at the first failing step."""
    steps = case.get("STEPS") or []
    if not steps:
        raise Skip("This test has no steps yet. Edit it and add at least one step.")
    for i, step in enumerate(steps, 1):
        label = STEP_LABELS.get(step.get("action"), step.get("action"))
        try:
            run_step(page, step)
        except Exception as exc:
            first = (str(exc).strip().splitlines() or [repr(exc)])[0][:260]
            raise AssertionError(f"Step {i} ({label}) failed: {first}")
    assert_no_db_error(page)
    return "PASS", f"All {len(steps)} step{'s' if len(steps) != 1 else ''} passed"


# Add each Excel TEST_KEY here only after its browser actions and assertions are implemented.
TESTS = {
    "HOURS_ABOVE_24": test_hours_above_24,
    "NEGATIVE_HOURS": test_negative_hours,
    "FUTURE_DATE_READONLY": test_future_date_readonly,
    "OLD_WEEK_READONLY": test_old_week_readonly,
    "OLD_WEEK_BOUNDARY_EDITABLE": test_old_week_boundary_editable,
    "WEEKEND_READONLY": test_weekend_readonly,
    "OVERTIME_REQUIRES_APPROVAL": test_overtime_requires_approval,
    "NORMAL_HOURS_APPROVED": test_normal_hours_approved,
    "ASSIGNED_ACTIVE_PROJECTS_ONLY": test_assigned_active_projects_only,
    "CROSS_EMPLOYEE_ACCESS_DENIED": test_cross_employee_access_denied,
    "WEEKLY_LIMIT": test_weekly_limit,
    "HOLIDAY_READONLY": test_holiday_readonly,
    "CUSTOM_STEPS": test_custom_steps,
}


# ----------------------------------------------------------------------------- Excel loading (with samples)
def sheet_rows(workbook, name):
    if name not in workbook.sheetnames:
        return []
    sheet = workbook[name]
    headers = [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() != "" for v in row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def load_cases(path=None):
    path = Path(path) if path else EXCEL_FILE
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Test_Cases" not in wb.sheetnames:
        raise ValueError("Excel workbook must contain a sheet named Test_Cases")
    for r in sheet_rows(wb, "Config"):
        k = str(r.get("KEY") or "").strip().upper()
        if k and r.get("VALUE") is not None:
            CONFIG[k] = str(r["VALUE"]).strip()
    parents = [r for r in sheet_rows(wb, "Test_Cases") if str(r.get("ENABLED", "")).strip().upper() == "Y"]
    extra = {}
    for r in sheet_rows(wb, "Samples"):
        extra.setdefault(str(r.get("TEST_CASE_ID") or "").strip(), []).append(r)
    cases = []
    for parent in parents:
        cid = str(parent.get("TEST_CASE_ID") or "").strip()
        if ONLY and cid.upper() not in ONLY:
            continue
        samples = []
        for v in split_samples(parent.get("TEST_DATA")):          # samples written as "25;30;99"
            s = dict(parent)
            s["TEST_DATA"] = v
            samples.append(s)
        for ex in extra.get(cid, []):                              # samples from the Samples sheet
            if str(ex.get("ENABLED", "Y")).strip().upper() == "N":
                continue
            s = dict(parent)
            for k, v in ex.items():
                if k in ("TEST_CASE_ID", "ENABLED") or v is None or str(v).strip() == "":
                    continue
                s[k] = v
            for v in split_samples(s.get("TEST_DATA")):
                s2 = dict(s)
                s2["TEST_DATA"] = v
                samples.append(s2)
        multi = len(samples) > 1
        for i, s in enumerate(samples, 1):
            s["_SAMPLE_NO"] = i
            s["_CASE_ID"] = f"{cid}-S{i}" if multi else cid
            label = s.get("SAMPLE_NAME") or (f"data={s['TEST_DATA']}" if s.get("TEST_DATA") is not None else f"sample {i}")
            s["_LABEL"] = f"{parent.get('TEST_NAME')} [{label}]" if multi and label else str(parent.get("TEST_NAME") or cid)
            cases.append(s)
    return cases


# ----------------------------------------------------------------------------- runner
def run_one(browser, storage_state, case):
    test_id = case["_CASE_ID"]
    test_name = case["_LABEL"]
    test_key = str(case.get("TEST_KEY") or "").strip().upper()
    expected = str(case.get("EXPECTED_RESULT") or "")
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", test_id + "_" + test_name).strip("_")[:120]
    video_dir = VIDEO_DIR / safe
    video_dir.mkdir(parents=True, exist_ok=True)
    started = time.time()
    status, detail, video_rel, shot_rel = "SKIP", "No Playwright test implemented for TEST_KEY: " + test_key, "", ""
    context = page = None
    ON_EVENT("start", test_id)
    if test_key not in TESTS:
        RESULTS.append(dict(test_case_id=test_id, name=test_name, test_key=test_key, expected=expected,
                            status=status, detail=detail, secs=0, video="", shot="", data=case.get("TEST_DATA")))
        print(f"[SKIP] {test_id} - {detail}")
        ON_EVENT("done", RESULTS[-1])
        return
    try:
        context = browser.new_context(storage_state=storage_state, viewport={"width": 1366, "height": 850},
                                      record_video_dir=str(video_dir), record_video_size={"width": 1366, "height": 850})
        page = context.new_page()
        page.set_default_timeout(20000)
        try:
            status, detail = TESTS[test_key](page, case)
        except Skip as s:
            status, detail = "SKIP", str(s)
        except Exception as exc:
            status, detail = "FAIL", (str(exc).strip().splitlines() or [repr(exc)])[0][:500]
        try:
            shot_path = OUT / f"{safe}.png"
            page.screenshot(path=str(shot_path), full_page=True, timeout=10000)
            shot_rel = str(shot_path.relative_to(OUT))
        except Exception:
            pass
    except Exception as exc:
        status, detail = "FAIL", (str(exc).strip().splitlines() or [repr(exc)])[0][:500]
    finally:
        if context:
            try:
                context.close()          # closing the context finalizes the .webm
                videos = [p for p in video_dir.glob("*.webm") if p.is_file()]
                if videos:
                    target = OUT / f"{safe}.webm"
                    shutil.move(str(videos[0]), str(target))
                    video_rel = str(target.relative_to(OUT))
            except Exception as exc:
                detail += " | Video warning: " + str(exc)[:150]
    elapsed = round(time.time() - started, 1)
    RESULTS.append(dict(test_case_id=test_id, name=test_name, test_key=test_key, expected=expected,
                        status=status, detail=detail, secs=elapsed, video=video_rel, shot=shot_rel,
                        data=case.get("TEST_DATA")))
    print(f"[{status}] {test_id} - {test_name} ({elapsed}s) {detail}")
    ON_EVENT("done", RESULTS[-1])


def write_report():
    csv_path = OUT / "results.csv"
    fields = ["test_case_id", "name", "test_key", "data", "expected", "status", "detail", "secs", "shot", "video"]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(RESULTS)
    counts = {s: sum(r["status"] == s for r in RESULTS) for s in ("PASS", "FAIL", "SKIP")}
    colors = {"PASS": "#166534", "FAIL": "#b91c1c", "SKIP": "#92400e"}
    rows = []
    for r in RESULTS:
        video = f'<a href="{html.escape(r["video"])}">Watch</a>' if r["video"] else "-"
        shot = f'<a href="{html.escape(r["shot"])}">Screenshot</a>' if r["shot"] else "-"
        rows.append("<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td style='color:{};font-weight:700'>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
            html.escape(r["test_case_id"]), html.escape(r["name"]), html.escape(str(r["data"] if r["data"] is not None else "")),
            html.escape(r["expected"]), colors.get(r["status"], "#000"), r["status"], html.escape(r["detail"]), shot, video))
    report = f'''<!doctype html><html><head><meta charset="utf-8"><title>Timesheet Test Results</title>
<style>body{{font-family:Arial,sans-serif;margin:24px}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ddd;padding:8px;text-align:left;vertical-align:top}}th{{background:#164e63;color:white}}tr:nth-child(even){{background:#f5f7f8}}</style></head><body>
<h2>Timesheet Automated Test Results</h2><p>Run: {datetime.now():%Y-%m-%d %H:%M:%S} | Total: {len(RESULTS)} | PASS: {counts['PASS']} | FAIL: {counts['FAIL']} | SKIP: {counts['SKIP']}</p>
<table><tr><th>ID</th><th>Test case</th><th>Sample data</th><th>Expected</th><th>Status</th><th>Actual / details</th><th>Screenshot</th><th>Video</th></tr>{''.join(rows)}</table></body></html>'''
    (OUT / "report.html").write_text(report, encoding="utf-8")
    print(f"\nPASS {counts['PASS']} | FAIL {counts['FAIL']} | SKIP {counts['SKIP']}")
    print(f"Report: {OUT / 'report.html'}\nCSV: {csv_path}\nEvidence folder: {OUT}")


def select_cases(excel_cases, only, extra_cases):
    """Excel cases filtered by the ids picked in the UI, followed by the UI-built cases."""
    if only is not None:
        wanted = {str(x).strip() for x in only}
        excel_cases = [c for c in excel_cases if c["_CASE_ID"] in wanted]
    return list(excel_cases) + list(extra_cases or [])


def run_job(excel_path, out_dir, base, cdp_url="http://127.0.0.1:9222", allow_save=False,
            headless=False, slow_mo=60, only=None, on_event=None, extra_cases=None):
    """Run the selected cases. Evidence goes to out_dir. Used by the web backend.
    only        = list of case ids (e.g. "TC-001-S2") picked in the UI; None = every enabled Excel case
    extra_cases = test cases built in the UI (already shaped like runner cases); excel_path may be None."""
    global BASE, CDP_URL, ALLOW_SAVE, HEADLESS, ONLY, OUT, VIDEO_DIR, RESULTS, CONFIG, SESSION, ON_EVENT
    BASE, CDP_URL, ALLOW_SAVE, HEADLESS = base.rstrip("/"), cdp_url, allow_save, headless
    ONLY = set()
    OUT = Path(out_dir)
    VIDEO_DIR = OUT / "videos"
    RESULTS, CONFIG, SESSION = [], {}, ""
    ON_EVENT = on_event or (lambda k, p: None)
    cases = select_cases(load_cases(excel_path) if excel_path else [], only, extra_cases)
    ON_EVENT("cases", [dict(id=c["_CASE_ID"], name=c["_LABEL"], key=str(c.get("TEST_KEY") or ""),
                            expected=str(c.get("EXPECTED_RESULT") or ""), data=c.get("TEST_DATA")) for c in cases])
    if not cases:
        raise ValueError("No test cases selected.")
    OUT.mkdir(parents=True, exist_ok=True)
    VIDEO_DIR.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        try:
            source_browser = p.chromium.connect_over_cdp(CDP_URL)
        except Exception as exc:
            raise RuntimeError(f"Could not connect to Chrome at {CDP_URL}. Start Chrome with --remote-debugging-port=9222 "
                               f"and log in to the app first. ({str(exc).splitlines()[0][:150]})")
        source_context = source_browser.contexts[0]
        app_part = BASE.split("/ords/")[-1].split("/")[0]
        for tab in source_context.pages:
            m = re.search(r"[?&]session=(\d+)", tab.url)
            if m and app_part in tab.url:
                SESSION = m.group(1)
                break
        storage_state = source_context.storage_state()
        browser = p.chromium.launch(channel="chrome", headless=headless, slow_mo=int(slow_mo))
        try:
            for case in cases:
                run_one(browser, storage_state, case)
        finally:
            browser.close()
    write_report()
    return RESULTS


def main():
    global SESSION
    if not BASE:
        print('Set APEX_BASE first, e.g. $env:APEX_BASE = "https://oracleapex.com/ords/r/abhi13_workspace/timesheet-app"')
        sys.exit(1)
    cases = load_cases()
    if not cases:
        print("No enabled test cases found in Excel (ENABLED must be Y).")
        return
    OUT.mkdir(parents=True, exist_ok=True)
    VIDEO_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Loaded {len(cases)} test run(s) from {EXCEL_FILE.name}. Save-tests {'ENABLED' if ALLOW_SAVE else 'disabled (set ALLOW_SAVE=1 to enable)'}.")
    with sync_playwright() as p:
        print("Connecting to your logged-in Chrome to copy the current APEX session...")
        source_browser = p.chromium.connect_over_cdp(CDP_URL)
        source_context = source_browser.contexts[0]
        app_part = BASE.split("/ords/")[-1].split("/")[0]
        for tab in source_context.pages:
            m = re.search(r"[?&]session=(\d+)", tab.url)
            if m and app_part in tab.url:
                SESSION = m.group(1)
                break
        if not SESSION:
            print("WARNING: no APEX session number found in open tabs. Open the app page in that Chrome and rerun if tests redirect to login.")
        storage_state = source_context.storage_state()
        print("Launching a separate controlled Chrome window (only the browser is recorded).")
        browser = p.chromium.launch(channel="chrome", headless=HEADLESS, slow_mo=int(os.environ.get("SLOW_MO", "60")))
        try:
            for case in cases:
                run_one(browser, storage_state, case)
        finally:
            browser.close()      # your original Chrome (CDP) is left open
    write_report()


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print("ERROR:", exc)
        sys.exit(1)
